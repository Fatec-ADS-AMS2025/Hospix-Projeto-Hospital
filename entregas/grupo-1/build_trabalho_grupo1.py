from __future__ import annotations

from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED
import shutil
import tempfile
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "entregas" / "grupo-1" / "Grupo 1 - Clean Code e SOLID Prescricao Medica.docx"


INK = "000000"
MUTED = "333333"
LIGHT = "F2F2F2"
CALLOUT = "FFFFFF"
CAUTION = "FFFFFF"
BLUE = INK
DARK_BLUE = INK


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_text(cell, text: str, bold: bool = False, color: str = INK) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(10)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP


def set_table_borders(table) -> None:
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), "D1D5DB")


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_sep)
    run._r.append(text)
    run._r.append(fld_end)


def configure_doc(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(3)
    section.left_margin = Cm(3)
    section.bottom_margin = Cm(2)
    section.right_margin = Cm(2)
    section.header_distance = Cm(1.25)
    section.footer_distance = Cm(1.25)
    section.different_first_page_header_footer = True

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(12)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.first_line_indent = Cm(1.25)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    for name, size, color, before, after in [
        ("Heading 1", 12, INK, 18, 6),
        ("Heading 2", 12, INK, 12, 6),
        ("Heading 3", 12, INK, 10, 4),
    ]:
        style = styles[name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)

    for name in ("List Bullet", "List Number"):
        style = styles[name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        style.font.size = Pt(12)
        style.font.color.rgb = RGBColor.from_string(INK)
        style.paragraph_format.line_spacing = 1.5

    if "CodeBlock" not in styles:
        code = styles.add_style("CodeBlock", 1)
    else:
        code = styles["CodeBlock"]
    code.font.name = "Consolas"
    code.font.size = Pt(9)
    code.font.color.rgb = RGBColor.from_string("111827")
    code.paragraph_format.left_indent = Cm(1.25)
    code.paragraph_format.right_indent = Cm(0)
    code.paragraph_format.first_line_indent = Cm(0)
    code.paragraph_format.space_after = Pt(1)
    code.paragraph_format.line_spacing = 1.0

    header = section.header.paragraphs[0]
    add_page_number(header)
    for run in header.runs:
        run.font.name = "Times New Roman"
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor.from_string(INK)


def add_title(doc: Document, text: str, subtitle: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_before = Pt(0)
    r = p.add_run("FACULDADE DE TECNOLOGIA")
    r.font.name = "Times New Roman"
    r.font.size = Pt(12)
    r.font.bold = True
    r.font.color.rgb = RGBColor.from_string(INK)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_before = Pt(0)
    r = p.add_run("CURSO DE ANALISE E DESENVOLVIMENTO DE SISTEMAS")
    r.font.name = "Times New Roman"
    r.font.size = Pt(12)
    r.font.bold = True
    r.font.color.rgb = RGBColor.from_string(INK)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_before = Pt(64)
    r = p.add_run("GRUPO 1 - QUALIDADE DO CODIGO")
    r.font.name = "Times New Roman"
    r.font.size = Pt(12)
    r.font.bold = True
    r.font.color.rgb = RGBColor.from_string(INK)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_before = Pt(110)
    r = p.add_run(text.upper())
    r.font.name = "Times New Roman"
    r.font.size = Pt(12)
    r.font.bold = True
    r.font.color.rgb = RGBColor.from_string(INK)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_before = Pt(0)
    r = p.add_run(subtitle)
    r.font.name = "Times New Roman"
    r.font.size = Pt(12)
    r.font.color.rgb = RGBColor.from_string(INK)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_before = Pt(150)
    r = p.add_run("SAO PAULO\n2026")
    r.font.name = "Times New Roman"
    r.font.size = Pt(12)
    r.font.color.rgb = RGBColor.from_string(INK)


def add_meta_table(doc: Document) -> None:
    doc.add_page_break()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_before = Pt(0)
    r = p.add_run("GRUPO 1 - QUALIDADE DO CODIGO")
    r.font.name = "Times New Roman"
    r.font.size = Pt(12)
    r.font.bold = True

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_before = Pt(80)
    r = p.add_run("CLEAN CODE E SOLID APLICADOS AO MODULO DE PRESCRICAO MEDICA")
    r.font.name = "Times New Roman"
    r.font.size = Pt(12)
    r.font.bold = True

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.left_indent = Cm(7)
    p.paragraph_format.space_before = Pt(110)
    p.paragraph_format.line_spacing = 1.0
    run = p.add_run(
        "Trabalho apresentado ao Seminario de Engenharia de Software, com foco no Grupo 1 "
        "- Qualidade do codigo, como parte da simulacao de desenvolvimento de um Sistema "
        "Hospitalar Inteligente."
    )
    run.font.name = "Times New Roman"
    run.font.size = Pt(12)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.left_indent = Cm(7)
    p.paragraph_format.line_spacing = 1.0
    run = p.add_run("Integrantes: preencher com os nomes dos integrantes do grupo.")
    run.font.name = "Times New Roman"
    run.font.size = Pt(12)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_before = Pt(96)
    r = p.add_run("SAO PAULO\n2026")
    r.font.name = "Times New Roman"
    r.font.size = Pt(12)


def heading(doc: Document, level: int, text: str) -> None:
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.left_indent = Cm(0)
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_before = Pt(18 if level == 1 else 12)
    p.paragraph_format.space_after = Pt(6)
    numeric = bool(text[:1].isdigit())
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT if numeric or level > 1 else WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text.upper() if level == 1 else text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(12)
    run.font.bold = True
    run.font.color.rgb = RGBColor.from_string(INK)


def para(doc: Document, text: str, style: str | None = None) -> None:
    p = doc.add_paragraph(style=style)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Cm(1.25)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.5
    p.add_run(text)


def bullet(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Cm(1.25)
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.5
    p.add_run(text)


def numbered(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.left_indent = Cm(1.25)
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.5
    p.add_run(text)


def code_block(doc: Document, title: str, code: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(title)
    r.font.name = "Times New Roman"
    r.font.size = Pt(10)
    r.font.bold = True
    r.font.color.rgb = RGBColor.from_string(DARK_BLUE)
    for line in code.strip("\n").splitlines():
        cp = doc.add_paragraph(style="CodeBlock")
        cp.paragraph_format.keep_together = False
        cp.add_run(line)


def add_callout(doc: Document, title: str, body: str, fill: str = CALLOUT) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Cm(1.25)
    p.paragraph_format.line_spacing = 1.5
    r = p.add_run(title)
    r.font.bold = True
    r.font.color.rgb = RGBColor.from_string(INK)
    r.font.name = "Times New Roman"
    r.font.size = Pt(12)
    p.add_run(": " + body)


def pretext_heading(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(18)
    p.paragraph_format.line_spacing = 1.5
    run = p.add_run(text.upper())
    run.font.name = "Times New Roman"
    run.font.size = Pt(12)
    run.font.bold = True
    run.font.color.rgb = RGBColor.from_string(INK)


def add_toc_field(doc: Document) -> None:
    pretext_heading(doc, "Sumario")
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Cm(0)
    run = p.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = 'TOC \\o "1-3" \\h \\z \\u'
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = "Atualizar sumario no Word."
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(separate)
    run._r.append(placeholder)
    run._r.append(end)


def force_black_text(docx_path: Path) -> None:
    with tempfile.TemporaryDirectory() as temp_dir:
        temp = Path(temp_dir)
        with ZipFile(docx_path, "r") as source:
            source.extractall(temp)

        for xml_path in (temp / "word").rglob("*.xml"):
            text = xml_path.read_text(encoding="utf-8")
            updated = re.sub(r"\s+w:themeColor=\"[^\"]*\"", "", text)
            updated = re.sub(r"\s+w:themeTint=\"[^\"]*\"", "", updated)
            updated = re.sub(r"\s+w:themeShade=\"[^\"]*\"", "", updated)
            updated = re.sub(r"<w:color\b[^>]*/>", '<w:color w:val="000000"/>', updated)
            if updated != text:
                xml_path.write_text(updated, encoding="utf-8")

        backup = docx_path.with_suffix(".docx.bak")
        shutil.copy2(docx_path, backup)
        with ZipFile(docx_path, "w", ZIP_DEFLATED) as target:
            for file_path in temp.rglob("*"):
                if file_path.is_file():
                    target.write(file_path, file_path.relative_to(temp).as_posix())
        backup.unlink(missing_ok=True)


def comparison_table(doc: Document) -> None:
    rows = [
        ("Responsabilidade", "Um HospitalService para pacientes, internacoes, prescricoes, exames, alta, alertas e faturamento.", "Um CreatePrescriptionUseCase focado em criar prescricao; dominio valida invariantes."),
        ("Dependencia", "Service depende diretamente de HospitalDbContext e SaveChanges.", "Application depende de IAdmissionRepository e IUnitOfWork; EF Core fica na Infrastructure."),
        ("Validacao", "Condicionais dentro do service; strings de status como 'INTERNADO' e 'ATIVA'.", "Value objects Dosage/DateRange e enums de dominio centralizam regras."),
        ("Testabilidade", "Para testar prescricao, o teste carrega muita infraestrutura do service geral.", "Use case pode ser testado com repositorio fake em memoria."),
        ("Evolucao", "Nova regra tende a aumentar o metodo e misturar fluxos.", "Novas regras entram no dominio, em use cases especificos ou por interfaces como IAlertRule."),
    ]
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    headers = ("Aspecto", "Legado", "Engineered")
    for cell, text in zip(table.rows[0].cells, headers):
        set_cell_shading(cell, LIGHT)
        set_cell_text(cell, text, bold=True, color=DARK_BLUE)
    for aspect, legacy, engineered in rows:
        cells = table.add_row().cells
        set_cell_text(cells[0], aspect, bold=True)
        set_cell_text(cells[1], legacy)
        set_cell_text(cells[2], engineered)


def criteria_table(doc: Document) -> None:
    rows = [
        ("Problemas de codigo ruim", "Metodo grande, acoplamento ao EF Core, strings de status, excesso de responsabilidade, baixa testabilidade."),
        ("Clean Code", "Nomes significativos, metodo curto, DTO tipado, classes menores, leitura do fluxo em alto nivel."),
        ("SOLID", "SRP em casos de uso, DIP por repositorios e Unit of Work, OCP em regras de alerta."),
        ("Refatoracao", "Comparacao direta entre HospitalService.CreatePrescription e CreatePrescriptionUseCase."),
        ("Ferramenta de analise", "dotnet format --verify-no-changes e testes automatizados de dominio/aplicacao."),
        ("Coerencia entre grupos", "Mesmo dominio hospitalar, mesmos termos: paciente, internacao, prescricao, exame, leito, fatura."),
    ]
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table)
    for cell, text in zip(table.rows[0].cells, ("Exigencia", "Como foi atendida")):
        set_cell_shading(cell, LIGHT)
        set_cell_text(cell, text, bold=True, color=DARK_BLUE)
    for key, value in rows:
        cells = table.add_row().cells
        set_cell_text(cells[0], key, bold=True)
        set_cell_text(cells[1], value)


def build() -> None:
    doc = Document()
    configure_doc(doc)

    add_title(
        doc,
        "Clean Code e SOLID Aplicados ao Modulo de Prescricao Medica",
        "Trabalho escrito - Grupo 1: Qualidade do codigo",
    )
    add_meta_table(doc)
    doc.add_page_break()

    pretext_heading(doc, "Resumo")
    para(
        doc,
        "Este trabalho apresenta a visao do Grupo 1 no seminario de Engenharia de Software: qualidade do codigo no modulo de prescricao medica de um sistema hospitalar web. O mesmo sistema usado pelos demais grupos possui pacientes, internacoes, prescricoes, exames, faturamento e comunicacao entre setores. O recorte deste grupo e menor e intencional: avaliar como a prescricao medica pode sair de um codigo legado funcional, porem dificil de manter, para uma implementacao mais legivel, testavel e preparada para evolucao.",
    )
    para(
        doc,
        "A comparacao usa dois projetos do mesmo repositorio. O hospital-legacy representa uma implementacao concentrada em um service amplo. O hospital-engineered representa a refatoracao com separacao entre API, Application, Domain e Infrastructure. O foco nao e defender que toda aplicacao pequena precisa de varias camadas, mas mostrar que, em um dominio hospitalar, regras clinicas e operacionais mudam com frequencia e precisam ficar visiveis no codigo.",
    )
    doc.add_page_break()
    add_toc_field(doc)
    doc.add_page_break()

    heading(doc, 1, "1. Introducao")
    para(
        doc,
        "Sistemas hospitalares nao podem ser avaliados apenas pela tela funcionando. Um cadastro de prescricao envolve medicamento, dose, frequencia, periodo, internacao ativa e efeitos em outros fluxos, como alta hospitalar e alertas. Quando essas regras ficam espalhadas ou escondidas em metodos grandes, cada alteracao passa a ter risco de quebrar comportamento clinico ou administrativo.",
    )
    para(
        doc,
        "O Tema 01 pede que o grupo demonstre como codigo ruim afeta manutencao, como aplicar Clean Code, como aplicar os principios SOLID usados no projeto, como refatorar codigo e como usar ao menos uma ferramenta de analise. Por isso, este documento trabalha com um antes e depois real: a prescricao criada dentro de HospitalService no sistema legado e a prescricao criada por CreatePrescriptionUseCase no sistema engineered.",
    )
    para(
        doc,
        "O objetivo tecnico e responder, com exemplos concretos, como escrever codigo sustentavel em um sistema hospitalar complexo. A resposta proposta e: separar responsabilidades, dar nomes que expressem o dominio, colocar regras invariantes no dominio, depender de abstracoes para persistencia e validar o comportamento com testes e ferramentas automaticas.",
    )
    heading(doc, 2, "Escopo fechado do Grupo 1")
    bullet(doc, "Cadastro de prescricao medica durante uma internacao ativa.")
    bullet(doc, "Inclusao de medicamento, dose, frequencia e periodo de administracao.")
    bullet(doc, "Validacao de dados obrigatorios e regra de periodo.")
    bullet(doc, "Separacao de responsabilidades entre controlador, caso de uso, dominio e persistencia.")
    bullet(doc, "Comparacao de codigo antes/depois com foco em qualidade, nao em regras de negocio completas.")
    heading(doc, 1, "2. Fundamentacao Teorica")
    heading(doc, 2, "2.1 Clean Code")
    para(
        doc,
        "Clean Code e uma abordagem de escrita de codigo voltada a clareza, manutencao e baixo custo de mudanca. Em vez de medir qualidade apenas por quantidade de linhas, a avaliacao considera se o proximo desenvolvedor consegue entender a intencao do codigo, localizar regras de negocio e alterar uma parte sem abrir um efeito colateral em outra.",
    )
    para(
        doc,
        "No modulo de prescricao, nomes como CreatePrescriptionUseCase, Dosage, DateRange, Admission e Prescription comunicam conceitos do hospital. Isso e melhor do que nomes genericos ou codigo que so revela sua intencao depois de ler varias condicionais. Um metodo curto tambem ajuda porque permite ler o fluxo principal sem misturar detalhes de persistencia, formatacao de strings e validacao de dominio.",
    )
    bullet(doc, "Nomes significativos: classes e metodos usam termos do dominio hospitalar.")
    bullet(doc, "Metodos curtos: cada metodo representa uma decisao ou operacao clara.")
    bullet(doc, "Organizacao: regras do dominio ficam no dominio; acesso a banco fica na infraestrutura.")
    bullet(doc, "Legibilidade: o fluxo principal aparece em alto nivel antes dos detalhes.")
    bullet(doc, "Padronizacao: excecoes de dominio, DTOs e interfaces seguem um mesmo estilo.")
    heading(doc, 2, "2.2 SOLID aplicado ao projeto")
    para(
        doc,
        "O Tema 01 pede apenas os principios aplicados ao projeto. Neste caso, os mais relevantes sao SRP, OCP e DIP. Eles foram usados de modo pratico, nao apenas teorico: cada um aparece em uma decisao de projeto do modulo de prescricao ou em funcionalidades proximas, como alertas relacionados a prescricoes ativas.",
    )
    numbered(doc, "SRP - Single Responsibility Principle: uma classe deve ter um motivo principal para mudar. A criacao de prescricao deixa de estar no service geral e passa para um caso de uso especifico.")
    numbered(doc, "OCP - Open/Closed Principle: o sistema deve permitir extensao sem modificar codigo central. As regras de alerta usam IAlertRule, permitindo novas regras sem alterar o orquestrador.")
    numbered(doc, "DIP - Dependency Inversion Principle: a camada de aplicacao depende de interfaces, nao de detalhes de banco. O caso de uso usa IAdmissionRepository e IUnitOfWork.")
    heading(doc, 1, "3. Problemas no Codigo Legado")
    para(
        doc,
        "No sistema legado, HospitalService concentra diversas responsabilidades: dashboard, pacientes, leitos, internacoes, alta, prescricoes, exames, alertas e faturamento. Mesmo que a prescricao em si tenha um trecho relativamente curto, ela esta dentro de uma classe que mistura muitas razoes para mudar. Isso viola SRP em nivel de classe e aumenta o custo de manutencao.",
    )
    code_block(
        doc,
        "Trecho do legado: hospital-legacy/backend/HospitalLegacy.Api/Services/HospitalService.cs",
        """
public PrescriptionDto CreatePrescription(CreatePrescriptionRequest request)
{
    if (string.IsNullOrWhiteSpace(request.MedicineName) || string.IsNullOrWhiteSpace(request.Dose))
    {
        throw new InvalidOperationException("Medicamento e dose sao obrigatorios.");
    }

    if (request.FrequencyHours <= 0)
    {
        throw new InvalidOperationException("Frequencia invalida.");
    }

    if (request.EndAt <= request.StartAt)
    {
        throw new InvalidOperationException("Fim da prescricao deve ser posterior ao inicio.");
    }

    var admission = db.Admissions.FirstOrDefault(x => x.Id == request.AdmissionId);
    if (admission == null || admission.Status != "INTERNADO")
    {
        throw new InvalidOperationException("Prescricao exige internacao ativa.");
    }

    db.Prescriptions.Add(prescription);
    db.SaveChanges();
    return ToPrescriptionDto(prescription);
}
""",
    )
    heading(doc, 2, "3.1 Diagnostico")
    bullet(doc, "Metodo dentro de um service grande, com varias responsabilidades fora do escopo de prescricao.")
    bullet(doc, "Acoplamento direto ao HospitalDbContext e ao EF Core, dificultando teste isolado.")
    bullet(doc, "Regras importantes escritas como condicionais soltas dentro do service.")
    bullet(doc, "Uso de strings de status, como 'INTERNADO' e 'ATIVA', sem protecao de tipo.")
    bullet(doc, "Persistencia, validacao, criacao de entidade e conversao para DTO no mesmo fluxo.")
    add_callout(
        doc,
        "Impacto pratico",
        "Em um hospital, uma regra de prescricao pode mudar por protocolo clinico, seguradora, auditoria ou seguranca do paciente. Se a regra esta escondida em um service geral, a equipe leva mais tempo para encontrar, testar e alterar o comportamento.",
        fill=CAUTION,
    )
    heading(doc, 1, "4. Refatoracao: Antes, Problemas e Depois")
    para(
        doc,
        "A refatoracao proposta nao muda a ideia funcional da prescricao. O medico continua criando uma prescricao para uma internacao ativa, informando medicamento, dose, frequencia e periodo. O que muda e a organizacao interna do codigo, para que a regra fique mais facil de entender, testar e evoluir.",
    )
    comparison_table(doc)
    heading(doc, 2, "4.1 Estrategia de refatoracao")
    numbered(doc, "Isolar o fluxo de criacao de prescricao em um caso de uso especifico.")
    numbered(doc, "Mover regras invariantes para o dominio: internacao ativa, dose obrigatoria, periodo valido e frequencia positiva.")
    numbered(doc, "Substituir dependencia direta do banco por interfaces da aplicacao.")
    numbered(doc, "Manter persistencia atras de Unit of Work.")
    numbered(doc, "Cobrir o comportamento principal com testes automatizados.")
    para(
        doc,
        "Essa estrategia melhora o design sem transformar o trabalho do Grupo 1 em um tema de arquitetura geral. A arquitetura aparece apenas como meio para qualidade de codigo: separa o que muda por regra de negocio, o que muda por infraestrutura e o que muda por entrada/saida da API.",
    )
    heading(doc, 1, "5. Codigo Depois: Caso de Uso de Prescricao")
    para(
        doc,
        "No sistema engineered, a criacao de prescricao e representada por CreatePrescriptionUseCase. O metodo ExecuteAsync mostra a intencao principal: buscar internacao, adicionar prescricao ao agregado, salvar a unidade de trabalho e retornar um DTO. A leitura fica direta porque os detalhes de validacao foram deslocados para objetos do dominio.",
    )
    code_block(
        doc,
        "Trecho refatorado: hospital-engineered/backend/src/Hospital.Application/UseCases/CreatePrescriptionUseCase.cs",
        """
public async Task<PrescriptionDto> ExecuteAsync(
    CreatePrescriptionRequest request,
    CancellationToken cancellationToken)
{
    var admission = await admissions.GetByIdAsync(request.AdmissionId, cancellationToken)
        ?? throw new DomainException("Internacao nao encontrada.");

    var prescription = admission.AddPrescription(
        Guid.NewGuid(),
        request.MedicineName,
        new Dosage(request.Dose),
        request.FrequencyHours,
        new DateRange(request.StartAt, request.EndAt));

    await unitOfWork.SaveChangesAsync(cancellationToken);
    return HospitalMapping.ToDto(prescription);
}
""",
    )
    heading(doc, 2, "5.1 Decisoes de Clean Code")
    bullet(doc, "O nome CreatePrescriptionUseCase deixa claro qual caso de uso e executado.")
    bullet(doc, "O metodo ExecuteAsync nao contem detalhes de EF Core.")
    bullet(doc, "O request e tipado por CreatePrescriptionRequest.")
    bullet(doc, "A criacao usa termos do dominio: Admission, Prescription, Dosage e DateRange.")
    bullet(doc, "A persistencia aparece como uma unica chamada: SaveChangesAsync.")
    heading(doc, 1, "6. Modelo de Dominio da Prescricao")
    para(
        doc,
        "A refatoracao fica mais forte porque o dominio protege regras importantes. O caso de uso nao precisa saber todos os detalhes de uma prescricao valida. Ele delega para Admission.AddPrescription e para os value objects. Isso reduz duplicacao e impede que outra parte do sistema crie uma prescricao ignorando validacoes essenciais.",
    )
    code_block(
        doc,
        "Trecho: Admission.AddPrescription",
        """
public Prescription AddPrescription(
    Guid prescriptionId,
    string medicineName,
    Dosage dose,
    int frequencyHours,
    DateRange period)
{
    EnsureActive("Prescricao exige internacao ativa.");
    var prescription = new Prescription(prescriptionId, Id, medicineName, dose, frequencyHours, period);
    Prescriptions.Add(prescription);
    return prescription;
}
""",
    )
    code_block(
        doc,
        "Trechos: Dosage e DateRange",
        """
public Dosage(string value)
{
    if (string.IsNullOrWhiteSpace(value))
        throw new DomainException("Dose e obrigatoria.");

    Value = value.Trim();
}

public DateRange(DateTime startAt, DateTime endAt)
{
    if (endAt <= startAt)
        throw new DomainException("Data final deve ser posterior a data inicial.");

    StartAt = startAt;
    EndAt = endAt;
}
""",
    )
    heading(doc, 2, "6.1 Por que separar Dose e Periodo?")
    para(
        doc,
        "Dose e periodo nao sao apenas strings ou datas soltas. Eles carregam regras. Quando Dosage rejeita vazio e DateRange rejeita fim anterior ao inicio, a validacao deixa de depender da memoria do programador. Qualquer codigo que tente criar uma prescricao passa pelas mesmas protecoes.",
    )
    heading(doc, 1, "7. SOLID na Pratica")
    heading(doc, 2, "7.1 SRP - Responsabilidade unica")
    para(
        doc,
        "O legado tem uma classe que muda por varios motivos: mudanca em faturamento, alertas, prescricao, alta, dashboard ou exames. No engineered, CreatePrescriptionUseCase muda quando muda o fluxo de criacao de prescricao. Prescription muda quando muda a regra da prescricao. Dosage muda quando muda a regra de dose. Essa divisao reduz o impacto de alteracoes.",
    )
    heading(doc, 2, "7.2 DIP - Dependencia de abstracoes")
    code_block(
        doc,
        "Interfaces usadas pelo caso de uso",
        """
public interface IAdmissionRepository
{
    Task<Admission?> GetByIdAsync(Guid id, CancellationToken cancellationToken);
    Task AddAsync(Admission admission, CancellationToken cancellationToken);
}

public interface IUnitOfWork
{
    Task SaveChangesAsync(CancellationToken cancellationToken);
}
""",
    )
    para(
        doc,
        "O caso de uso nao conhece EF Core, tabela ou DbSet. Ele conhece o contrato necessario para executar a operacao. Isso facilita teste com fake repository e permite trocar detalhes de infraestrutura sem reescrever regra de aplicacao.",
    )
    heading(doc, 2, "7.3 OCP - Extensao por regras")
    code_block(
        doc,
        "Exemplo proximo ao modulo: regras de alerta por IAlertRule",
        """
public interface IAlertRule
{
    IEnumerable<AlertDto> Evaluate(HospitalAlertContext context);
}

public class ActivePrescriptionAlertRule : IAlertRule
{
    public IEnumerable<AlertDto> Evaluate(HospitalAlertContext context)
    {
        var active = context.Prescriptions.Count(x => x.Status == PrescriptionStatus.Active);
        if (active > 0)
            yield return new AlertDto("info", "Prescricoes ativas", "...", "prescricoes");
    }
}
""",
    )
    heading(doc, 1, "8. Diagramas e Modelagem do Fluxo")
    heading(doc, 2, "8.1 Fluxo legado")
    code_block(
        doc,
        "Visao simplificada",
        """
HospitalController
    -> HospitalService
        -> valida request com condicionais
        -> consulta HospitalDbContext diretamente
        -> cria PrescriptionRecord
        -> salva com db.SaveChanges()
        -> retorna PrescriptionDto
""",
    )
    para(
        doc,
        "Nesse fluxo, a mesma classe conhece HTTP indiretamente, regra de negocio, banco de dados e formato de resposta. A modelagem fica simples de desenhar, mas cara de manter, porque quase tudo aponta para o service central.",
    )
    heading(doc, 2, "8.2 Fluxo engineered")
    code_block(
        doc,
        "Visao simplificada",
        """
PrescriptionsController
    -> CreatePrescriptionUseCase
        -> IAdmissionRepository.GetByIdAsync()
        -> Admission.AddPrescription()
            -> Prescription
            -> Dosage
            -> DateRange
        -> IUnitOfWork.SaveChangesAsync()
        -> HospitalMapping.ToDto()
""",
    )
    para(
        doc,
        "Nesse fluxo, a API recebe a requisicao, a camada de aplicacao orquestra o caso de uso e o dominio protege as regras essenciais. O repositorio fica separado para que a camada de aplicacao nao dependa diretamente do banco de dados.",
    )
    add_callout(
        doc,
        "Leitura para a apresentacao",
        "O diagrama engineered parece maior porque mostra responsabilidades separadas. A vantagem e que cada caixa tem um motivo claro para existir; no legado, a simplicidade visual esconde acoplamento.",
    )
    heading(doc, 1, "9. Ferramentas de Analise e Evidencias")
    para(
        doc,
        "O Tema 01 exige demonstrar ao menos uma ferramenta de analise. Neste trabalho foram usadas ferramentas do proprio ecossistema .NET, coerentes com a implementacao em C#/.NET: dotnet format e testes automatizados. O dotnet format --verify-no-changes verifica padronizacao de estilo sem alterar arquivos. Os testes validam comportamento de dominio e aplicacao.",
    )
    code_block(
        doc,
        "Comandos executados",
        """
dotnet format .\\Hospital.Engineered.slnx --verify-no-changes --verbosity minimal
dotnet test .\\backend\\tests\\Hospital.Domain.Tests\\Hospital.Domain.Tests.csproj
dotnet test .\\backend\\tests\\Hospital.Application.Tests\\Hospital.Application.Tests.csproj
dotnet test .\\tests\\HospitalLegacy.Api.Tests\\HospitalLegacy.Api.Tests.csproj --no-build
""",
    )
    heading(doc, 2, "9.1 Resultado registrado")
    bullet(doc, "dotnet format no engineered: concluido com codigo de saida 0, sem alteracoes exigidas.")
    bullet(doc, "Hospital.Domain.Tests: 4 testes aprovados, 0 falhas.")
    bullet(doc, "Hospital.Application.Tests: 3 testes aprovados, 0 falhas.")
    bullet(doc, "HospitalLegacy.Api.Tests com --no-build: 2 testes aprovados, 0 falhas.")
    heading(doc, 2, "9.2 Teste do caso de uso")
    para(
        doc,
        "O teste principal do caso de uso cria uma prescricao valida para uma internacao ativa, verifica o medicamento retornado, confirma o status ATIVA e garante que a unidade de trabalho foi salva. Isso demonstra que a refatoracao nao ficou apenas estetica: o comportamento essencial do modulo continua verificavel por teste automatizado.",
    )
    heading(doc, 1, "10. Discussao")
    para(
        doc,
        "A solucao engineered e mais eficiente do ponto de vista de engenharia porque reduz custo de mudanca e melhora testabilidade. E importante separar eficiencia de desempenho bruto. Para um cadastro de prescricao, a diferenca de tempo de execucao entre os dois desenhos nao e o ponto principal. O ganho real esta em manter regras clinicas compreensiveis, isoladas e verificaveis.",
    )
    para(
        doc,
        "No legado, uma alteracao simples como 'bloquear prescricao fora do horario permitido' provavelmente entraria dentro de HospitalService. Com o tempo, esse service ficaria maior e mais sensivel. No engineered, a mesma regra poderia ir para o dominio, para um value object ou para um caso de uso especifico, mantendo o fluxo de aplicacao legivel.",
    )
    para(
        doc,
        "Tambem existe trade-off. O engineered tem mais arquivos, interfaces e classes. Para um exercicio pequeno isso pode parecer excesso. Porem, o contexto do seminario e um sistema hospitalar com pacientes, internacoes, prescricoes, exames, faturamento e comunicacao entre setores. Nesse tipo de sistema, o custo de uma regra mal posicionada cresce rapidamente. Por isso, a separacao e coerente com a complexidade proposta na atividade.",
    )
    heading(doc, 2, "Integracao com os outros grupos")
    bullet(doc, "Grupo 1 explica qualidade do codigo e refatoracao no modulo de prescricao.")
    bullet(doc, "Grupo 2 pode aprofundar regras de negocio: alta bloqueada, internacao ativa, faturamento e validacoes clinicas.")
    bullet(doc, "Grupo 3 pode aprofundar arquitetura: camadas, DDD, repositorios, gateway de faturamento e separacao entre servicos.")
    para(
        doc,
        "A continuidade entre apresentacoes fica natural: o Grupo 1 mostra por que o codigo precisa ser sustentavel; o Grupo 2 mostra quais regras sustentam o dominio; o Grupo 3 mostra como a arquitetura organiza esses elementos em escala.",
    )
    heading(doc, 1, "11. Conclusao")
    para(
        doc,
        "Codigo sustentavel em sistemas hospitalares complexos nasce da combinacao entre clareza, responsabilidade bem definida e validacao automatizada. No modulo de prescricao, isso significa evitar que regras clinicas fiquem perdidas em um service geral, representar conceitos importantes com nomes do dominio e proteger invariantes em classes adequadas.",
    )
    para(
        doc,
        "A refatoracao do legado para o engineered mostra essa mudanca. O codigo antes funcionava, mas concentrava validacao, acesso a dados, criacao de registro e persistencia no mesmo service. O codigo depois distribui as responsabilidades: o controller recebe HTTP, o use case orquestra, o dominio valida e a infraestrutura persiste. Essa organizacao esta coerente com o objetivo do Tema 01 e com o tema integrador do sistema hospitalar.",
    )
    add_callout(
        doc,
        "Resposta final",
        "Escrevemos codigo sustentavel em sistemas hospitalares complexos quando as regras importantes ficam explicitas no dominio, os fluxos de aplicacao ficam curtos, as dependencias apontam para abstracoes e o comportamento e validado por testes e ferramentas automaticas.",
    )
    heading(doc, 1, "Referencias")
    references = [
        "MARTIN, Robert C. Clean Code: A Handbook of Agile Software Craftsmanship. Prentice Hall, 2008.",
        "MARTIN, Robert C. Agile Software Development, Principles, Patterns, and Practices. Prentice Hall, 2002.",
        "FOWLER, Martin. Refactoring: Improving the Design of Existing Code. Addison-Wesley, 2018.",
        "MICROSOFT. .NET code analysis overview. Disponivel em: https://learn.microsoft.com/dotnet/fundamentals/code-analysis/overview",
        "MICROSOFT. dotnet format command. Disponivel em: https://learn.microsoft.com/dotnet/core/tools/dotnet-format",
        "MICROSOFT. Dependency injection in .NET. Disponivel em: https://learn.microsoft.com/dotnet/core/extensions/dependency-injection",
        "Repositorio do projeto Hospix: exemplos em hospital-legacy e hospital-engineered, consultados em 4 de junho de 2026.",
    ]
    for ref in references:
        bullet(doc, ref)

    doc.save(OUT)
    force_black_text(OUT)


if __name__ == "__main__":
    build()
